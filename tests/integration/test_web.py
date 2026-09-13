"""Exercise the actual local API, CLI subprocesses and delivery archive."""

import io
import time
import zipfile
from pathlib import Path

import pytest

pytest.importorskip("fastapi")
from fastapi.testclient import TestClient

from build_skills.web.server import create_app

HEADERS = {"X-Build-Skills": "local"}
ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture
def client(tmp_path: Path) -> TestClient:
    return TestClient(create_app(ROOT / "examples/local-files/config.toml", tmp_path))


def wait(client: TestClient, identifier: str) -> dict:
    deadline = time.monotonic() + 20
    while time.monotonic() < deadline:
        result = client.get(f"/api/jobs/{identifier}").json()
        if not result["busy"]:
            return result
        time.sleep(0.05)
    pytest.fail("Workflow did not settle")


def test_upload_approve_deliver(client: TestClient) -> None:
    response = client.post(
        "/api/materials", headers=HEADERS, files=[("files", ("notes.md", b"Copy text exactly."))]
    )
    assert response.status_code == 200, response.text
    material = response.json()
    options = client.get("/api/options").json()
    settings = {
        key: options[key]
        for key in ("name", "goal", "roles", "models", "max_rounds", "repetitions", "minimum_score")
    }
    settings["material"] = material["id"]
    settings["providers"] = {
        name: {"model": p["model"], "reasoning_effort": p["reasoning_effort"]}
        for name, p in options["providers"].items()
    }
    response = client.post("/api/jobs", headers=HEADERS, json=settings)
    assert response.status_code == 200, response.text
    identifier = response.json()["id"]
    state = wait(client, identifier)
    assert state["status"] == "awaiting_approval", state
    assert client.get(f"/api/jobs/{identifier}/download").status_code == 409
    bad = client.post(
        f"/api/jobs/{identifier}/approve",
        headers=HEADERS,
        json={"brief": state["brief"], "digest": "outdated"},
    )
    assert bad.status_code == 409
    response = client.post(
        f"/api/jobs/{identifier}/approve",
        headers=HEADERS,
        json={"brief": state["brief"], "digest": state["brief_digest"]},
    )
    assert response.status_code == 200, response.text
    state = wait(client, identifier)
    assert state["status"] == "delivered", state
    archive = client.get(f"/api/jobs/{identifier}/download")
    assert archive.status_code == 200, archive.text
    with zipfile.ZipFile(io.BytesIO(archive.content)) as bundle:
        assert set(bundle.namelist()) == {"skill/SKILL.md", "report.json"}
    assert client.get(f"/api/jobs/{identifier}/report").json()
    assert client.get("/api/jobs").json()[0]["id"] == identifier


def test_local_boundary_and_bad_upload(client: TestClient) -> None:
    assert client.post("/api/materials").status_code == 403
    assert client.get("/api/options", headers={"host": "evil.example"}).status_code == 403
    assert (
        client.post(
            "/api/materials",
            headers=HEADERS | {"origin": "https://evil.example"},
            files={"files": ("a.txt", b"hello")},
        ).status_code
        == 403
    )
    for name, content in [
        ("../notes.md", b"hello"),
        ("empty.txt", b""),
        ("bad.pdf", b"not a pdf"),
        ("data.bin", b"binary"),
    ]:
        response = client.post("/api/materials", headers=HEADERS, files={"files": (name, content)})
        assert response.status_code == 400, response.text
    assert client.get("/api/jobs/not-an-id").status_code == 404


def test_docx_conversion(client: TestClient) -> None:
    from docx import Document

    document = Document()
    document.add_paragraph("Synthetic instructions")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Expected output"
    stream = io.BytesIO()
    document.save(stream)
    response = client.post(
        "/api/materials", headers=HEADERS, files={"files": ("sample.docx", stream.getvalue())}
    )
    assert response.status_code == 200, response.text
    assert response.json()["files"][0]["text"] == "Synthetic instructions\nExpected output"
